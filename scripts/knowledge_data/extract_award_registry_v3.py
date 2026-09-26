"""荣誉台账全量提取（2026-09-26 用户要求"知识清单每一条都详细录入"）。

源：三、学校荣誉及办学成果资料【必需】/1、学校荣誉/6.荣誉获得时间（全校）/
（2019-2020 至 2025-2026 各学年获奖登记 xls/xlsx + 科组补充 doc）+ 近三年荣誉.pptx。

xls→xlrd、xlsx→openpyxl、doc→Word COM、pptx→PowerPoint COM。
产出 docs_text_v3/荣誉台账/<文件名>.txt（逐 sheet、逐行 | 拼接）；
stdout 汇总字符数。只落地文本，不进 git（师生姓名 PII）。
"""
import sys
from pathlib import Path

import xlrd

sys.stdout.reconfigure(encoding="utf-8")
SRC = Path(r"D:/SRP/AI_school/数据清单/数据清单/三、学校荣誉及办学成果资料【必需】/1、学校荣誉")
OUT = Path(__file__).resolve().parent / "docs_text_v3" / "荣誉台账"
SKIP = ("~$", ".DS_Store", "Thumbs.db")


def sheet_text_xls(path: Path) -> str:
    wb = xlrd.open_workbook(str(path))
    parts = []
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            vals = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
            if any(vals):
                rows.append(" | ".join(vals))
        parts.append(f"--- sheet: {sh.name} ({sh.nrows}行) ---\n" + "\n".join(rows))
    return "\n\n".join(parts)


def sheet_text_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sh in wb.worksheets:
        rows = []
        for row in sh.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in row]
            if any(vals):
                rows.append(" | ".join(vals))
        parts.append(f"--- sheet: {sh.title} ({len(rows)}行) ---\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def com_doc_text(word, path: Path) -> str:
    doc = word.Documents.Open(str(path), ReadOnly=True)
    text = doc.Content.Text.replace("\r", "\n").replace("\x07", "\n")
    tables = []
    for i, tbl in enumerate(doc.Tables, 1):
        cells = []
        for r in range(1, tbl.Rows.Count + 1):
            row = []
            for c in range(1, tbl.Columns.Count + 1):
                try:
                    row.append(tbl.Cell(r, c).Range.Text.replace("\r\x07", "").strip())
                except Exception:
                    row.append("")
            cells.append(" | ".join(row))
        tables.append(f"--- 表格{i} ---\n" + "\n".join(cells))
    doc.Close(False)
    body = text.strip()
    if tables:
        body = (body + "\n\n" if body else "") + "\n\n".join(tables)
    return body


def com_ppt_text(ppt, path: Path) -> str:
    pres = ppt.Presentations.Open(str(path), ReadOnly=True, WithWindow=False)
    parts = []
    for slide in pres.Slides:
        lines = []
        for shape in slide.Shapes:
            if shape.HasTextFrame and shape.TextFrame.HasText:
                t = shape.TextFrame.TextRange.Text.strip()
                if t:
                    lines.append(t)
        parts.append(f"--- 第{slide.SlideIndex}页 ---\n" + "\n".join(lines))
    pres.Close()
    return "\n\n".join(parts)


def main():
    word = None
    ppt = None
    jobs = []
    OUT.mkdir(parents=True, exist_ok=True)
    for sub in sorted(SRC.iterdir()):
        if not sub.is_dir() or sub.name == "6.荣誉获得时间（全校）":
            continue
    reg = SRC / "6.荣誉获得时间（全校）"
    for p in sorted(reg.rglob("*")):
        if not p.is_file() or p.name.startswith(SKIP):
            continue
        if p.suffix.lower() in (".xls", ".xlsx", ".doc", ".docx"):
            jobs.append(p)
    extra = [SRC / "近三年荣誉.pptx"]
    jobs.extend(p for p in extra if p.exists())

    index = []
    for p in jobs:
        ext = p.suffix.lower()
        rel = p.relative_to(SRC)
        dst = OUT / p.name.replace("/", "_")
        dst = dst.with_suffix(".txt")
        note = ""
        try:
            if ext == ".xls":
                text = sheet_text_xls(p)
            elif ext == ".xlsx":
                text = sheet_text_xlsx(p)
            elif ext == ".pptx":
                if ppt is None:
                    import win32com.client

                    ppt = win32com.client.Dispatch("PowerPoint.Application")
                text = com_ppt_text(ppt, p)
            else:
                if word is None:
                    import win32com.client

                    word = win32com.client.Dispatch("Word.Application")
                    word.Visible = False
                text = com_doc_text(word, p)
        except Exception as ex:
            note = f"EXTRACT_FAIL: {ex}"
            text = ""
        dst.write_text(
            f"【来源】{rel}\n【提取】荣誉台账全量（2026-09-26）\n\n{text}", encoding="utf-8")
        index.append({"path": str(rel), "chars": len(text), "note": note})
        print(f"{len(text):>7}  {rel}{' <<<' if note else ''}")

    (OUT / "_index.json").write_text(
        __import__("json").dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n共 {len(index)} 个台账 → {OUT}")


if __name__ == "__main__":
    main()
