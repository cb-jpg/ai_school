"""
Document processing for knowledge base entries.
Handles file parsing, text extraction, and chunking.
"""
import os
import re
import asyncio
from pathlib import Path
from typing import List, Optional, Tuple
from datetime import datetime
from loguru import logger

try:
    import aiohttp
    from bs4 import BeautifulSoup
    WEBSUPPORT = True
except ImportError:
    WEBSUPPORT = False
    logger.warning("aiohttp or beautifulsoup4 not installed. Web scraping disabled.")

from .models import KnowledgeEntry, Chunk, SourceType, KnowledgeCategory, KnowledgeStatus


class DocumentProcessor:
    """Process documents for knowledge base ingestion"""

    def __init__(self, knowledge_dir: str = "data/knowledge"):
        self.knowledge_dir = Path(knowledge_dir)
        self.documents_dir = self.knowledge_dir / "documents"
        self.processed_dir = self.knowledge_dir / "processed"

        # Ensure directories exist
        self.documents_dir.mkdir(parents=True, exist_ok=True)
        self.processed_dir.mkdir(parents=True, exist_ok=True)

        # Chunk size for text splitting
        self.chunk_size = 500
        self.chunk_overlap = 50

    async def process_file(
        self,
        file_path: str,
        title: str,
        category: KnowledgeCategory,
        tags: List[str],
        created_by: Optional[str] = None
    ) -> Tuple[KnowledgeEntry, List[Chunk]]:
        """
        Process an uploaded file and create knowledge entry with chunks.

        Args:
            file_path: Path to the uploaded file
            title: Title for the knowledge entry
            category: Knowledge category
            tags: List of tags
            created_by: User who created the entry

        Returns:
            Tuple of (KnowledgeEntry, List[Chunk])
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_ext = file_path.suffix.lower()
        text_content = ""

        try:
            # Extract text based on file type
            if file_ext == '.txt':
                text_content = await self._process_txt(file_path)
            elif file_ext == '.md':
                text_content = await self._process_markdown(file_path)
            elif file_ext == '.pdf':
                text_content = await self._process_pdf(file_path)
            elif file_ext in ['.doc', '.docx']:
                text_content = await self._process_word(file_path)
            elif file_ext in ['.xls', '.xlsx']:
                text_content = await self._process_excel(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_ext}")

            if not text_content or len(text_content.strip()) < 10:
                raise ValueError("Extracted text is too short or empty")

            # Create chunks
            chunks = self._chunk_text(text_content)

            # Create knowledge entry
            entry = KnowledgeEntry(
                title=title,
                category=category,
                tags=tags,
                source_type=SourceType.FILE,
                file_path=str(file_path),
                file_name=file_path.name,
                status=KnowledgeStatus.INDEXED,
                chunk_count=len(chunks),
                summary=self._generate_summary(text_content),
                created_by=created_by
            )

            # Save processed data
            await self._save_processed_data(entry.id, text_content, chunks)

            logger.info(f"Successfully processed file: {file_path.name}, chunks: {len(chunks)}")
            return entry, chunks

        except Exception as e:
            logger.error(f"Error processing file {file_path}: {e}")
            # Return entry with error status
            return KnowledgeEntry(
                title=title,
                category=category,
                tags=tags,
                source_type=SourceType.FILE,
                file_path=str(file_path),
                file_name=file_path.name,
                status=KnowledgeStatus.ERROR,
                error_message=str(e),
                chunk_count=0,
                created_by=created_by
            ), []

    async def process_url(
        self,
        url: str,
        title: Optional[str],
        category: KnowledgeCategory,
        tags: List[str],
        created_by: Optional[str] = None
    ) -> Tuple[KnowledgeEntry, List[Chunk]]:
        """
        Process a URL and create knowledge entry with chunks.

        Args:
            url: URL to fetch and process
            title: Optional title (will use page title if not provided)
            category: Knowledge category
            tags: List of tags
            created_by: User who created the entry

        Returns:
            Tuple of (KnowledgeEntry, List[Chunk])
        """
        if not WEBSUPPORT:
            raise RuntimeError("Web scraping not available. Install aiohttp and beautifulsoup4.")

        try:
            async with aiohttp.ClientSession() as session:
                async with session.get(url, timeout=aiohttp.ClientTimeout(total=30)) as response:
                    if response.status != 200:
                        raise ValueError(f"Failed to fetch URL: {response.status}")
                    html = await response.text()

            # Parse HTML
            soup = BeautifulSoup(html, 'html.parser')

            # Extract title if not provided
            if not title:
                title_tag = soup.find('title')
                title = title_tag.text.strip() if title_tag else url

            # Extract main content
            text_content = self._extract_main_content(soup)

            if not text_content or len(text_content.strip()) < 10:
                raise ValueError("Extracted web content is too short or empty")

            # Create chunks
            chunks = self._chunk_text(text_content)

            # Create knowledge entry
            entry = KnowledgeEntry(
                title=title or url,
                category=category,
                tags=tags,
                source_type=SourceType.URL,
                source_url=url,
                status=KnowledgeStatus.INDEXED,
                chunk_count=len(chunks),
                summary=self._generate_summary(text_content),
                created_by=created_by
            )

            # Save processed data
            await self._save_processed_data(entry.id, text_content, chunks)

            logger.info(f"Successfully processed URL: {url}, chunks: {len(chunks)}")
            return entry, chunks

        except Exception as e:
            logger.error(f"Error processing URL {url}: {e}")
            # Return entry with error status
            return KnowledgeEntry(
                title=title or url,
                category=category,
                tags=tags,
                source_type=SourceType.URL,
                source_url=url,
                status=KnowledgeStatus.ERROR,
                error_message=str(e),
                chunk_count=0,
                created_by=created_by
            ), []

    async def create_manual_entry(
        self,
        title: str,
        content: str,
        category: KnowledgeCategory,
        tags: List[str],
        summary: Optional[str] = None,
        created_by: Optional[str] = None
    ) -> Tuple[KnowledgeEntry, List[Chunk]]:
        """
        Create a manual knowledge entry from provided text.

        Args:
            title: Title for the knowledge entry
            content: Text content
            category: Knowledge category
            tags: List of tags
            summary: Optional summary
            created_by: User who created the entry

        Returns:
            Tuple of (KnowledgeEntry, List[Chunk])
        """
        if not content or len(content.strip()) < 10:
            raise ValueError("Content is too short or empty")

        try:
            # Create chunks
            chunks = self._chunk_text(content)

            # Create knowledge entry
            entry = KnowledgeEntry(
                title=title,
                category=category,
                tags=tags,
                source_type=SourceType.MANUAL,
                status=KnowledgeStatus.INDEXED,
                chunk_count=len(chunks),
                summary=summary or self._generate_summary(content),
                created_by=created_by
            )

            # Save processed data
            await self._save_processed_data(entry.id, content, chunks)

            logger.info(f"Successfully created manual entry: {title}, chunks: {len(chunks)}")
            return entry, chunks

        except Exception as e:
            logger.error(f"Error creating manual entry {title}: {e}")
            raise

    async def _process_txt(self, file_path: Path) -> str:
        """Process plain text file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

    async def _process_markdown(self, file_path: Path) -> str:
        """Process markdown file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        # Remove markdown syntax for cleaner text
        content = re.sub(r'[#*`_\[\]]', '', content)
        return content

    async def _process_pdf(self, file_path: Path) -> str:
        """Process PDF file"""
        try:
            import PyPDF2
            text = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text.append(page.extract_text())
            return '\n'.join(text)
        except ImportError:
            raise RuntimeError("PyPDF2 not installed. Install with: pip install PyPDF2")

    async def _process_word(self, file_path: Path) -> str:
        """Process Word document"""
        try:
            import docx
            doc = docx.Document(file_path)
            text = []
            for paragraph in doc.paragraphs:
                text.append(paragraph.text)
            return '\n'.join(text)
        except ImportError:
            raise RuntimeError("python-docx not installed. Install with: pip install python-docx")

    async def _process_excel(self, file_path: Path) -> str:
        """Process Excel file"""
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(file_path)
            text = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = ' '.join(str(cell) if cell is not None else '' for cell in row)
                    if row_text.strip():
                        text.append(row_text)
            return '\n'.join(text)
        except ImportError:
            raise RuntimeError("openpyxl not installed. Install with: pip install openpyxl")

    def _extract_main_content(self, soup) -> str:
        """Extract main content from HTML"""
        # Remove script and style elements
        for element in soup(['script', 'style', 'nav', 'footer', 'header']):
            element.decompose()

        # Try to find main content area
        main_content = soup.find('main') or soup.find('article') or soup.find('div', class_=re.compile('content|main', re.I))

        if main_content:
            text = main_content.get_text(separator='\n', strip=True)
        else:
            text = soup.get_text(separator='\n', strip=True)

        # Clean up whitespace
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        return '\n'.join(lines)

    def _chunk_text(self, text: str) -> List[Chunk]:
        """
        Split text into chunks for vectorization.

        Args:
            text: Text to chunk

        Returns:
            List of Chunk objects
        """
        # Clean text
        text = re.sub(r'\s+', ' ', text)

        chunks = []
        start = 0
        chunk_index = 0

        while start < len(text):
            end = start + self.chunk_size

            # Try to break at sentence boundary
            if end < len(text):
                # Look for sentence endings
                for delimiter in ['。', '！', '？', '.', '!', '?', '\n']:
                    last_delimiter = text.rfind(delimiter, start, end)
                    if last_delimiter != -1:
                        end = last_delimiter + 1
                        break

            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append(Chunk(
                    content=chunk_text,
                    source_id="",  # Will be set when entry is created
                    chunk_index=chunk_index,
                    metadata={"start": start, "end": end}
                ))
                chunk_index += 1

            start = end - self.chunk_overlap

        return chunks

    def _generate_summary(self, text: str, max_length: int = 200) -> str:
        """Generate a summary from text"""
        # Take first meaningful sentences
        sentences = re.split(r'[。！？.!?]', text)
        summary_parts = []
        total_length = 0

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            if total_length + len(sentence) > max_length:
                break
            summary_parts.append(sentence)
            total_length += len(sentence) + 1

        summary = '。'.join(summary_parts)
        if len(summary) < len(text) and len(text) > max_length:
            summary += '...'
        return summary

    async def _save_processed_data(self, entry_id: str, content: str, chunks: List[Chunk]):
        """Save processed data to disk"""
        entry_dir = self.processed_dir / entry_id
        entry_dir.mkdir(exist_ok=True)

        # Save original content
        content_file = entry_dir / "content.txt"
        with open(content_file, 'w', encoding='utf-8') as f:
            f.write(content)

        # Save chunks
        chunks_file = entry_dir / "chunks.json"
        import json
        with open(chunks_file, 'w', encoding='utf-8') as f:
            json.dump([chunk.model_dump() for chunk in chunks], f, ensure_ascii=False, indent=2)

        logger.debug(f"Saved processed data for entry {entry_id}")

    async def load_chunks(self, entry_id: str) -> List[Chunk]:
        """Load chunks for an entry"""
        chunks_file = self.processed_dir / entry_id / "chunks.json"
        if not chunks_file.exists():
            return []

        import json
        with open(chunks_file, 'r', encoding='utf-8') as f:
            data = json.load(f)

        return [Chunk(**chunk) for chunk in data]
